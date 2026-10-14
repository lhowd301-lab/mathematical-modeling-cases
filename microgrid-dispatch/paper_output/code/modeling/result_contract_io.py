from __future__ import annotations

import hashlib
import json
import shutil
from datetime import datetime
from pathlib import Path

import numpy as np
import openpyxl

from common_storage_lp import aggregate_four_hour


ROOT = Path(__file__).resolve().parents[3]
TEMPLATE_DIR = ROOT / "problem_files" / "附件5"
RESULT_DIR = ROOT / "paper_output" / "results"
TABLE_DIR = ROOT / "paper_output" / "tables"


def sha256_file(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_json(path, payload):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )


def _copy_template(template_name, target):
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(TEMPLATE_DIR / template_name, target)
    return openpyxl.load_workbook(target)


def _write_charge_sheet(worksheet, day_results):
    worksheet.delete_rows(2, worksheet.max_row)
    row = 2
    for day in sorted(day_results):
        result = day_results[day]
        charge = aggregate_four_hour(result["charge_kwh"])
        discharge = aggregate_four_hour(result["discharge_kwh"])
        for index, label in enumerate(
            ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
        ):
            worksheet.cell(row=row + index, column=1, value=day if index == 0 else None)
            worksheet.cell(row=row + index, column=2, value=label)
            worksheet.cell(row=row + index, column=3, value=float(charge[index]))
            worksheet.cell(row=row + index, column=4, value=float(discharge[index]))
        worksheet.cell(row=row, column=5, value="00:00")
        worksheet.cell(row=row, column=6, value=float(result["initial_soc_kwh"]))
        worksheet.cell(row=row + 1, column=5, value="24:00")
        worksheet.cell(row=row + 1, column=6, value=float(result["terminal_soc_kwh"]))
        row += 6


def _write_purchase_sheet(worksheet, day_results):
    day_results = {str(key): value for key, value in day_results.items()}
    row_lookup = {}
    for row in range(2, worksheet.max_row + 1):
        raw_date = worksheet.cell(row=row, column=1).value
        if raw_date is None:
            continue
        if isinstance(raw_date, datetime):
            date_text = raw_date.date().isoformat()
        else:
            date_text = str(raw_date).split(" ")[0]
        row_lookup[date_text] = row
    for day, result in sorted(day_results.items()):
        row = row_lookup.get(day)
        if row is None:
            continue
        for index, value in enumerate(result["purchase_kwh"], start=2):
            worksheet.cell(row=row, column=index, value=float(value))
        worksheet.cell(row=row, column=146, value=float(np.sum(result["purchase_kwh"])))
        worksheet.cell(row=row, column=147, value=float(result["cost_yuan"]))


def _merge_emergency(day, result):
    events = []
    values = np.asarray(result["emergency_kwh"], dtype=float)
    starts = result.get("slot_start", [])
    ends = result.get("slot_end", [])
    current_start = None
    current_end = None
    current_value = 0.0
    for index, value in enumerate(values):
        if value > 1e-6 and current_start is None:
            current_start = starts[index]
            current_end = ends[index]
            current_value = float(value)
        elif value > 1e-6 and current_start is not None:
            current_end = ends[index]
            current_value += float(value)
        elif value <= 1e-6 and current_start is not None:
            events.append((day, current_start, current_end, current_value))
            current_start = None
            current_end = None
            current_value = 0.0
    if current_start is not None:
        events.append((day, current_start, current_end, current_value))
    return events


def _write_emergency_sheet(worksheet, day_results):
    worksheet.delete_rows(2, worksheet.max_row)
    row = 2
    for day in sorted(day_results):
        events = _merge_emergency(day, day_results[day])
        for index, (event_day, start, end, value) in enumerate(events):
            worksheet.cell(row=row, column=1, value=event_day if index == 0 else None)
            worksheet.cell(row=row, column=2, value="{} {}".format(start, end))
            worksheet.cell(row=row, column=3, value=value)
            row += 1
    if row == 2:
        worksheet.cell(row=2, column=1, value="无")
        worksheet.cell(row=2, column=2, value="本年度未触发紧急购电")
        worksheet.cell(row=2, column=3, value=0.0)


def write_result1(path, result):
    target = Path(path)
    workbook = _copy_template("result1.xlsx", target)
    plan = workbook["计划购电量"]
    for index, value in enumerate(result["purchase_kwh"], start=2):
        plan.cell(row=index, column=2, value=float(value))

    charge_sheet = workbook["充放电量"]
    charge = aggregate_four_hour(result["charge_kwh"])
    discharge = aggregate_four_hour(result["discharge_kwh"])
    for index in range(6):
        charge_sheet.cell(row=2 + index, column=2, value=float(charge[index]))
        charge_sheet.cell(row=2 + index, column=3, value=float(discharge[index]))
    charge_sheet.cell(row=2, column=5, value=float(result["initial_soc_kwh"]))
    charge_sheet.cell(row=3, column=5, value=float(result["terminal_soc_kwh"]))
    workbook.save(target)


def write_result2(path, result):
    target = Path(path)
    workbook = _copy_template("result2.xlsx", target)
    _write_purchase_sheet(workbook["计划购电量"], result["days"])
    _write_charge_sheet(workbook["充放电量"], result["days"])
    _write_emergency_sheet(workbook["紧急购电量"], result["days"])
    workbook.save(target)


def write_result3(path, result):
    target = Path(path)
    workbook = _copy_template("result3.xlsx", target)
    _write_purchase_sheet(
        workbook["计划购电量"],
        {
            day: {
                **value,
                "purchase_kwh": value["initial_plan_kwh"],
                "cost_yuan": value["initial_plan_cost_yuan"],
            }
            for day, value in result["days"].items()
        },
    )
    _write_purchase_sheet(workbook["调整购电量"], result["days"])
    _write_charge_sheet(workbook["充放电量"], result["days"])
    _write_emergency_sheet(workbook["紧急购电量"], result["days"])
    workbook.save(target)


def provenance(run_context, question_id, output_artifacts):
    normalized_artifacts = []
    for artifact in output_artifacts:
        artifact_path = Path(artifact)
        try:
            normalized = artifact_path.resolve().relative_to(ROOT).as_posix()
        except Exception:
            normalized = artifact_path.as_posix()
        normalized_artifacts.append(normalized)
    return {
        "source_code_path": run_context["source_code_path"],
        "source_code_sha256": run_context["source_code_sha256"],
        "run_command": run_context["run_command"],
        "run_exit_code": int(run_context["run_exit_code"]),
        "question_id": question_id,
        "output_artifacts": normalized_artifacts,
    }


def table_record(question_id, table_id, path, title, description):
    path = Path(path)
    return {
        "question_id": question_id,
        "table_id": table_id,
        "title": title,
        "description": description,
        "path": path.relative_to(ROOT).as_posix(),
        "exists": path.exists(),
        "bytes": path.stat().st_size if path.exists() else 0,
        "sha256": sha256_file(path) if path.exists() else "",
    }


def build_run_manifest(
    status,
    question_ids,
    code_paths,
    input_paths,
    output_paths,
    run_command,
    started_at,
    finished_at,
    notes=None,
):
    def file_record(path):
        item = Path(path)
        return {
            "path": item.relative_to(ROOT).as_posix(),
            "exists": item.exists(),
            "bytes": item.stat().st_size if item.exists() else 0,
            "sha256": sha256_file(item) if item.exists() else "",
        }

    script_path = Path(code_paths[0])
    run_record = {
        "script": script_path.relative_to(ROOT).as_posix(),
        "script_sha256": sha256_file(script_path),
        "returncode": 0,
        "question_ids": list(question_ids),
        "input_files": [file_record(path) for path in input_paths],
        "output_artifacts": [file_record(path) for path in output_paths],
        "started_at": started_at,
        "finished_at": finished_at,
    }
    return {
        "schema_version": "1.0",
        "generated_by": "paper_output/code/modeling/run_modeling.py",
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "status": status,
        "runs": [run_record],
        "question_ids": list(question_ids),
        "run_command": run_command,
        "working_directory": ROOT.as_posix(),
        "started_at": started_at,
        "finished_at": finished_at,
        "code": [file_record(path) for path in code_paths],
        "inputs": [file_record(path) for path in input_paths],
        "outputs": [file_record(path) for path in output_paths],
        "notes": list(notes or []),
    }
