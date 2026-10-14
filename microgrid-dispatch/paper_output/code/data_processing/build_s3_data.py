from __future__ import annotations

import hashlib
import json
import math
from collections import Counter
from datetime import date, datetime, time, timedelta
from pathlib import Path

import matplotlib
import numpy as np
import openpyxl
import pandas as pd


matplotlib.use("Agg")
import matplotlib.pyplot as plt


ROOT = Path(__file__).resolve().parents[3]
PROBLEM_DIR = ROOT / "problem_files"
OUTPUT_DIR = ROOT / "paper_output"
CLEAN_DIR = OUTPUT_DIR / "data_cleaned"
PLAN_DIR = OUTPUT_DIR / "plan"
FIGURE_DIR = OUTPUT_DIR / "figures"
EDA_FIGURE_DIR = FIGURE_DIR / "eda"
MANIFEST_FILE = OUTPUT_DIR / "input_manifest.json"
MODEL_ROUTE_FILE = PLAN_DIR / "model_route.json"
RAW_LOAD_REPORT_FILE = CLEAN_DIR / "load_report.json"

ATTACHMENT1 = PROBLEM_DIR / "附件1.xlsx"
ATTACHMENT2 = PROBLEM_DIR / "附件2.xlsx"
ATTACHMENT3 = PROBLEM_DIR / "附件3.xlsx"
ATTACHMENT4 = PROBLEM_DIR / "附件4.xlsx"
TEMPLATE_DIR = PROBLEM_DIR / "附件5"

FIRST_DATE = date(2025, 1, 1)
LAST_DATE = date(2025, 12, 31)
OUTPUT_START = date(2025, 2, 1)
DELTA_HOURS = 1.0 / 6.0
ORIGIN_HOURS = [0, 6, 12, 18]
SPECIFIED_DATES = [
    date(2025, 3, 20),
    date(2025, 6, 21),
    date(2025, 9, 23),
    date(2025, 12, 21),
]


def configure_plot_style():
    plt.rcParams["font.sans-serif"] = [
        "Microsoft YaHei",
        "SimHei",
        "Source Han Sans SC",
        "DejaVu Sans",
    ]
    plt.rcParams["axes.unicode_minus"] = False
    plt.rcParams["figure.dpi"] = 120
    plt.rcParams["savefig.dpi"] = 300
    plt.rcParams["axes.grid"] = True
    plt.rcParams["grid.alpha"] = 0.25
    plt.rcParams["axes.spines.top"] = False
    plt.rcParams["axes.spines.right"] = False


def sha256_file(path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_json(path):
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def write_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def normalize_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    parsed = pd.to_datetime(text, errors="raise")
    return parsed.date()


def time_text(value):
    if isinstance(value, datetime):
        return value.strftime("%H:%M")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    text = str(value).strip()
    if text.endswith("+1"):
        text = text[:-2].strip()
    if " " in text:
        text = text.split()[-1]
    parts = text.split(":")
    if len(parts) >= 2:
        return "{:02d}:{:02d}".format(int(parts[0]), int(parts[1]))
    return text


def parse_interval_start(day, label):
    plus_one = "+1" in str(label)
    text = time_text(label)
    hour, minute = [int(item) for item in text.split(":")[:2]]
    result = datetime.combine(day, time(hour=hour, minute=minute))
    if plus_one:
        result += timedelta(days=1)
    return result


def read_sheet_rows(path, sheet_name=None):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
    rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    workbook.close()
    return rows


def numeric(value):
    if value is None or value == "":
        return None
    return float(value)


def load_attachment1():
    rows = read_sheet_rows(ATTACHMENT1)
    header = rows[0]
    data = rows[1:]
    return header, data


def load_attachment2():
    load_rows = read_sheet_rows(ATTACHMENT2, "小区负载")
    pv_rows = read_sheet_rows(ATTACHMENT2, "光伏发电实际功率")
    return load_rows, pv_rows


def load_attachment3():
    return read_sheet_rows(ATTACHMENT3)


def load_attachment4():
    return read_sheet_rows(ATTACHMENT4)


def load_result_templates():
    contracts = {}
    for name in ["result1.xlsx", "result2.xlsx", "result3.xlsx", "result4-2.xlsx", "result4-3.xlsx"]:
        workbook = openpyxl.load_workbook(TEMPLATE_DIR / name, read_only=True, data_only=True)
        contracts[name] = {}
        for worksheet in workbook.worksheets:
            rows = list(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=min(worksheet.max_row, 3),
                    values_only=True,
                )
            )
            contracts[name][worksheet.title] = {
                "rows": worksheet.max_row,
                "cols": worksheet.max_column,
                "headers": [[None if value is None else str(value) for value in row] for row in rows],
            }
        workbook.close()
    return contracts


def build_interval_metadata(attachment1_rows, result1_labels):
    metadata = []
    for index, row in enumerate(attachment1_rows, start=1):
        raw_label = row[0]
        start = parse_interval_start(FIRST_DATE, raw_label)
        metadata.append(
            {
                "interval_id": index,
                "attachment_time_label": time_text(raw_label),
                "result1_interval_label": str(result1_labels[index - 1]),
                "start": start,
                "end": start + timedelta(minutes=10),
                "target_hour_index": int(
                    ((start - datetime.combine(FIRST_DATE, time(0, 0))).total_seconds()) // 3600
                ),
            }
        )
    return metadata


def build_attachment1_table(attachment1_rows, interval_metadata, result2_labels):
    records = []
    for index, row in enumerate(attachment1_rows):
        meta = interval_metadata[index]
        price = numeric(row[1])
        load = numeric(row[2])
        pv = numeric(row[3])
        records.append(
            {
                "interval_id": meta["interval_id"],
                "attachment_time_label": meta["attachment_time_label"],
                "result1_interval_label": meta["result1_interval_label"],
                "result2_interval_label": str(result2_labels[index]),
                "slot_start": meta["start"].isoformat(sep=" "),
                "slot_end": meta["end"].isoformat(sep=" "),
                "price_fixed_yuan_per_kwh": price,
                "load_kw": load,
                "pv_forecast_kw": pv,
                "load_kwh": load * DELTA_HOURS,
                "pv_forecast_kwh": pv * DELTA_HOURS,
            }
        )
    return pd.DataFrame(records)


def result2_interval_labels():
    rows = read_sheet_rows(TEMPLATE_DIR / "result2.xlsx", "计划购电量")
    header = rows[0]
    labels = header[1:145]
    if len(labels) != 144:
        raise ValueError("result2 template does not contain 144 interval columns")
    return labels


def result1_interval_labels():
    rows = read_sheet_rows(TEMPLATE_DIR / "result1.xlsx", "计划购电量")
    labels = [row[0] for row in rows[1:145]]
    if len(labels) != 144:
        raise ValueError("result1 template does not contain 144 interval rows")
    return labels


def build_annual_table(load_rows, pv_rows, price_rows, attachment1_table, interval_metadata, result2_labels):
    load_header = load_rows[0]
    pv_header = pv_rows[0]
    price_header = price_rows[0]
    if len(load_header) != 145 or len(pv_header) != 145 or len(price_header) != 145:
        raise ValueError("annual attachments must contain one date column and 144 data columns")
    load_data = load_rows[1:]
    pv_data = pv_rows[1:]
    price_data = price_rows[1:]
    if not (len(load_data) == len(pv_data) == len(price_data) == 365):
        raise ValueError("annual attachments must contain 365 data rows")

    fixed_prices = attachment1_table["price_fixed_yuan_per_kwh"].tolist()
    records = []
    for day_index, (load_row, pv_row, price_row) in enumerate(zip(load_data, pv_data, price_data)):
        day = normalize_date(load_row[0])
        if day != normalize_date(pv_row[0]) or day != normalize_date(price_row[0]):
            raise ValueError("annual attachment dates do not align at row {}".format(day_index + 2))
        for interval_index, meta in enumerate(interval_metadata):
            load_kw = numeric(load_row[interval_index + 1])
            pv_kw = numeric(pv_row[interval_index + 1])
            price = numeric(price_row[interval_index + 1])
            records.append(
                {
                    "date": day.isoformat(),
                    "day_index": day_index + 1,
                    "interval_id": meta["interval_id"],
                    "attachment_time_label": meta["attachment_time_label"],
                    "result_template_interval_label": str(result2_labels[interval_index]),
                    "slot_start": datetime.combine(day, meta["start"].time()).isoformat(sep=" ")
                    if meta["start"].date() == FIRST_DATE
                    else (datetime.combine(day + timedelta(days=1), time(0, 0))).isoformat(sep=" "),
                    "slot_end": (
                        datetime.combine(day, meta["end"].time())
                        if meta["end"].date() == FIRST_DATE
                        else datetime.combine(day + timedelta(days=1), time(0, 10))
                    ).isoformat(sep=" "),
                    "target_hour_index": meta["target_hour_index"],
                    "load_kw": load_kw,
                    "pv_actual_kw": pv_kw,
                    "price_actual_yuan_per_kwh": price,
                    "price_fixed_yuan_per_kwh": fixed_prices[interval_index],
                    "load_kwh": load_kw * DELTA_HOURS,
                    "pv_actual_kwh": pv_kw * DELTA_HOURS,
                    "is_official_output_period": day >= OUTPUT_START,
                }
            )
    return pd.DataFrame(records)


def build_forecast_long(forecast_rows):
    records = []
    current_date = None
    date_blank_count = 0
    for row_index, row in enumerate(forecast_rows[1:], start=2):
        raw_date = row[0]
        if raw_date is None or str(raw_date).strip() == "":
            date_blank_count += 1
        else:
            current_date = normalize_date(raw_date)
        if current_date is None:
            raise ValueError("first forecast row has no date")
        origin_hour = int(time_text(row[1]).split(":")[0])
        origin_dt = datetime.combine(current_date, time(origin_hour, 0))
        for horizon in range(1, 25):
            value = numeric(row[1 + horizon])
            target = origin_dt + timedelta(hours=horizon)
            target_hour_index = int(
                (
                    target
                    - datetime.combine(current_date, time(0, 0))
                ).total_seconds()
                // 3600
            )
            records.append(
                {
                    "forecast_date": current_date.isoformat(),
                    "forecast_origin_hour": origin_hour,
                    "forecast_origin_time": "{:02d}:00".format(origin_hour),
                    "forecast_origin_datetime": origin_dt.isoformat(sep=" "),
                    "forecast_horizon_hour": horizon,
                    "source_column": header_name_for_horizon(horizon),
                    "target_datetime": target.isoformat(sep=" "),
                    "target_date": target.date().isoformat(),
                    "target_hour_index": target_hour_index,
                    "pv_forecast_kw": value,
                }
            )
    return pd.DataFrame(records), date_blank_count


def header_name_for_horizon(horizon):
    return "预报{}小时".format(horizon)


def build_forecast_by_origin(annual_table, forecast_long):
    lookup = {}
    for row in forecast_long.itertuples(index=False):
        key = (row.forecast_date, int(row.forecast_origin_hour), int(row.forecast_horizon_hour))
        lookup[key] = row.pv_forecast_kw

    records = []
    base_by_day = {}
    for row in annual_table.itertuples(index=False):
        base_by_day.setdefault(row.date, []).append(row)

    for day in sorted(base_by_day):
        intervals = sorted(base_by_day[day], key=lambda item: item.interval_id)
        for origin_hour in ORIGIN_HOURS:
            origin_dt = datetime.combine(normalize_date(day), time(origin_hour, 0))
            for item in intervals:
                slot_start = datetime.fromisoformat(item.slot_start)
                hour_index = int(item.target_hour_index)
                difference = hour_index - origin_hour
                if slot_start < origin_dt:
                    status = "past_locked"
                elif 1 <= difference <= 24:
                    status = "available"
                elif hour_index == 0 and origin_hour == 0:
                    status = "night_zero_fallback"
                else:
                    status = "current_hour_hold"
                value = lookup.get((day, origin_hour, difference)) if 1 <= difference <= 24 else None
                if status == "night_zero_fallback":
                    value = 0.0
                records.append(
                    {
                        "forecast_date": day,
                        "forecast_origin_hour": origin_hour,
                        "forecast_origin_time": "{:02d}:00".format(origin_hour),
                        "forecast_origin_datetime": origin_dt.isoformat(sep=" "),
                        "interval_id": int(item.interval_id),
                        "slot_start": item.slot_start,
                        "slot_end": item.slot_end,
                        "target_hour_index": hour_index,
                        "forecast_horizon_hour": difference if 1 <= difference <= 24 else None,
                        "source_column": header_name_for_horizon(difference) if 1 <= difference <= 24 else "",
                        "pv_forecast_kw": value,
                        "pv_forecast_kwh": value * DELTA_HOURS if value is not None else None,
                        "forecast_status": status,
                    }
                )
    return pd.DataFrame(records)


def build_forecast_asof(annual_table, forecast_long):
    lookup = {}
    for row in forecast_long.itertuples(index=False):
        key = (
            normalize_date(row.forecast_date),
            int(row.forecast_origin_hour),
            int(row.forecast_horizon_hour),
        )
        lookup[key] = row.pv_forecast_kw

    records = []
    for item in annual_table.itertuples(index=False):
        day = normalize_date(item.date)
        slot_start = datetime.fromisoformat(item.slot_start)
        target_hour = int(item.target_hour_index)
        candidates = []
        previous_day = day - timedelta(days=1)
        previous_horizon = target_hour + 6
        if previous_day >= FIRST_DATE and 1 <= previous_horizon <= 24:
            candidates.append((previous_day, 18, previous_horizon))
        for origin_hour in ORIGIN_HOURS:
            origin_dt = datetime.combine(day, time(origin_hour, 0))
            horizon = target_hour - origin_hour
            if origin_dt < slot_start and 1 <= horizon <= 24:
                candidates.append((day, origin_hour, horizon))

        if candidates:
            candidates.sort(
                key=lambda value: (
                    value[0],
                    value[1],
                )
            )
            source_date, source_origin, horizon = candidates[-1]
            source_date = normalize_date(source_date)
            source_dt = datetime.combine(
                source_date,
                time(source_origin, 0),
            )
            value = lookup[(source_date, source_origin, horizon)]
            status = "available"
        else:
            source_date = None
            source_origin = None
            horizon = None
            source_dt = None
            value = 0.0
            status = "outside_dataset_night_zero_fallback"

        records.append(
            {
                "date": day.isoformat(),
                "interval_id": int(item.interval_id),
                "slot_start": item.slot_start,
                "slot_end": item.slot_end,
                "target_hour_index": target_hour,
                "forecast_source_date": source_date.isoformat() if source_date else "",
                "forecast_source_origin_hour": source_origin,
                "forecast_source_origin_time": "{:02d}:00".format(source_origin) if source_origin is not None else "",
                "forecast_source_datetime": source_dt.isoformat(sep=" ") if source_dt else "",
                "forecast_horizon_hour": horizon,
                "source_column": header_name_for_horizon(horizon) if horizon else "",
                "pv_forecast_kw": value,
                "pv_forecast_kwh": value * DELTA_HOURS,
                "forecast_status": status,
            }
        )
    return pd.DataFrame(records)


def dataframe_profile(frame, path):
    missing = {column: int(frame[column].isna().sum()) for column in frame.columns}
    return {
        "path": path.relative_to(ROOT).as_posix(),
        "rows": int(len(frame)),
        "columns": [str(column) for column in frame.columns],
        "missing_counts": missing,
        "sha256": sha256_file(path),
    }


def numeric_summary(series):
    clean = pd.to_numeric(series, errors="coerce").dropna()
    return {
        "min": float(clean.min()),
        "max": float(clean.max()),
        "mean": float(clean.mean()),
        "std": float(clean.std(ddof=0)),
    }


def build_monthly_stats(annual_table):
    frame = annual_table.copy()
    frame["month"] = pd.to_datetime(frame["date"]).dt.strftime("%Y-%m")
    grouped = frame.groupby("month", as_index=False).agg(
        mean_load_kw=("load_kw", "mean"),
        mean_pv_actual_kw=("pv_actual_kw", "mean"),
        mean_price_yuan_per_kwh=("price_actual_yuan_per_kwh", "mean"),
        total_load_kwh=("load_kwh", "sum"),
        total_pv_actual_kwh=("pv_actual_kwh", "sum"),
    )
    return grouped


def build_forecast_availability(forecast_by_origin, forecast_asof):
    by_origin = (
        forecast_by_origin.groupby(["forecast_origin_hour", "forecast_status"], as_index=False)
        .size()
        .rename(columns={"size": "row_count"})
    )
    asof = (
        forecast_asof.groupby(["forecast_status"], as_index=False)
        .size()
        .rename(columns={"size": "row_count"})
    )
    asof["forecast_origin_hour"] = "asof"
    return pd.concat([by_origin, asof], ignore_index=True, sort=False)


def write_eda_figures(annual_table, forecast_asof, monthly_stats):
    EDA_FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    figures = []

    fig1_path = EDA_FIGURE_DIR / "fig_data_monthly_profile.png"
    fig, axes = plt.subplots(3, 1, figsize=(10, 9), sharex=True)
    axes[0].plot(monthly_stats["month"], monthly_stats["mean_load_kw"], color="#1f77b4", marker="o")
    axes[0].set_ylabel("月均负载 (kW)")
    axes[0].set_title("2025 年月度平均负载、光伏与电价")
    axes[1].plot(monthly_stats["month"], monthly_stats["mean_pv_actual_kw"], color="#2ca02c", marker="o")
    axes[1].set_ylabel("月均光伏 (kW)")
    axes[2].plot(
        monthly_stats["month"],
        monthly_stats["mean_price_yuan_per_kwh"],
        color="#d62728",
        marker="o",
    )
    axes[2].set_ylabel("月均电价 (元/kWh)")
    axes[2].set_xlabel("月份")
    axes[2].tick_params(axis="x", rotation=45)
    fig.tight_layout()
    fig.savefig(fig1_path, bbox_inches="tight")
    plt.close(fig)
    figures.append(
        {
            "figure_id": "fig_data_1",
            "question_id": "DATA",
            "title": "2025 年月度平均负载、光伏与电价",
            "chart_type": "monthly_profile",
            "data_source": "paper_output/data_cleaned/annual_10min.csv",
            "output_path": fig1_path.relative_to(ROOT).as_posix(),
            "paper_usage": "数据说明与预处理",
            "purpose": "展示全年负载、光伏和电价的季节变化。",
            "status": "ready",
        }
    )

    fig2_path = EDA_FIGURE_DIR / "fig_data_sample_days.png"
    fig, axes = plt.subplots(3, 1, figsize=(12, 10), sharex=True)
    colors = ["#1f77b4", "#ff7f0e", "#2ca02c", "#d62728"]
    for color, selected in zip(colors, SPECIFIED_DATES):
        day_frame = annual_table[annual_table["date"] == selected.isoformat()].sort_values("interval_id")
        if day_frame.empty:
            continue
        x = day_frame["slot_start"].str.slice(11, 16)
        axes[0].plot(x, day_frame["load_kw"], color=color, label=selected.isoformat())
        axes[1].plot(x, day_frame["pv_actual_kw"], color=color, label=selected.isoformat())
        axes[2].plot(x, day_frame["price_actual_yuan_per_kwh"], color=color, label=selected.isoformat())
    axes[0].set_ylabel("负载 (kW)")
    axes[0].set_title("四个指定日期的负载、光伏实际值与波动电价")
    axes[1].set_ylabel("光伏实际值 (kW)")
    axes[2].set_ylabel("电价 (元/kWh)")
    axes[2].set_xlabel("时段起点")
    axes[0].legend(loc="best", ncol=2, fontsize=8)
    for axis in axes:
        axis.tick_params(axis="x", rotation=60)
        for index, label in enumerate(axis.get_xticklabels()):
            if index % 12 != 0:
                label.set_visible(False)
    fig.tight_layout()
    fig.savefig(fig2_path, bbox_inches="tight")
    plt.close(fig)
    figures.append(
        {
            "figure_id": "fig_data_2",
            "question_id": "DATA",
            "title": "四个指定日期的负载、光伏实际值与波动电价",
            "chart_type": "sample_day_profiles",
            "data_source": "paper_output/data_cleaned/annual_10min.csv",
            "output_path": fig2_path.relative_to(ROOT).as_posix(),
            "paper_usage": "问题二与问题四数据说明",
            "purpose": "核验指定日期的输入特征和日内变化。",
            "status": "ready",
        }
    )

    fig3_path = EDA_FIGURE_DIR / "fig_data_pv_actual_vs_asof_forecast.png"
    merged = annual_table[["date", "interval_id", "pv_actual_kw"]].merge(
        forecast_asof[["date", "interval_id", "pv_forecast_kw"]],
        on=["date", "interval_id"],
        how="left",
    )
    merged["month"] = pd.to_datetime(merged["date"]).dt.strftime("%Y-%m")
    monthly = merged.groupby("month", as_index=False).agg(
        mean_actual=("pv_actual_kw", "mean"),
        mean_forecast=("pv_forecast_kw", "mean"),
    )
    fig, axis = plt.subplots(figsize=(10, 4.8))
    axis.plot(monthly["month"], monthly["mean_actual"], color="#2ca02c", marker="o", label="实际光伏")
    axis.plot(
        monthly["month"],
        monthly["mean_forecast"],
        color="#9467bd",
        marker="s",
        label="最新可用预报",
    )
    axis.set_title("月度平均实际光伏与最新可用预报")
    axis.set_ylabel("平均功率 (kW)")
    axis.set_xlabel("月份")
    axis.tick_params(axis="x", rotation=45)
    axis.legend()
    fig.tight_layout()
    fig.savefig(fig3_path, bbox_inches="tight")
    plt.close(fig)
    figures.append(
        {
            "figure_id": "fig_data_3",
            "question_id": "DATA",
            "title": "月度平均实际光伏与最新可用预报",
            "chart_type": "forecast_alignment",
            "data_source": "paper_output/data_cleaned/annual_10min.csv; paper_output/data_cleaned/pv_forecast_10min_asof.csv",
            "output_path": fig3_path.relative_to(ROOT).as_posix(),
            "paper_usage": "问题三数据说明与预报映射验证",
            "purpose": "检查预报展开后与年度实际光伏的总体对齐情况。",
            "status": "ready",
        }
    )

    return figures


def build_planned_figures(model_route):
    planned = []
    questions = model_route.get("questions") or model_route.get("question_routes") or []
    for question in questions:
        qid = str(question.get("question_id"))
        for item in question.get("figures") or []:
            path = str(item.get("expected_path") or "")
            planned.append(
                {
                    "figure_id": item.get("figure_id"),
                    "question_id": qid,
                    "title": item.get("title"),
                    "chart_type": "model_result",
                    "data_source": "pending_model_results",
                    "output_path": path,
                    "paper_usage": "结果分析",
                    "purpose": item.get("purpose"),
                    "status": "ready" if path and (ROOT / path).exists() else "planned",
                }
            )
    return planned


def build_data_plan(annual_table, attachment1_table, forecast_long, forecast_by_origin, forecast_asof, templates):
    return {
        "schema_version": "1.1",
        "stage": "S3",
        "status": "complete",
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "generated_by": "paper_output/code/data_processing/build_s3_data.py",
        "source_contracts": [
            "paper_output/input_manifest.json",
            "paper_output/step1/problem_analysis.json",
            "paper_output/plan/model_route.json",
        ],
        "canonical_time_contract": {
            "interval_count": 144,
            "attachment_time_is_interval_start": True,
            "first_interval": "00:10-00:20",
            "last_interval": "次日 00:00-00:10",
            "result1_first_label": "0:10-0:20",
            "result1_last_label": "0:00+1-0:10+1",
            "note": "附件时间标签、result1 行标签和 result2 列标签在位置上一一对应。",
        },
        "datasets": [
            {
                "path": "paper_output/data_cleaned/attachment1_10min.csv",
                "rows": int(len(attachment1_table)),
                "fields": list(attachment1_table.columns),
                "used_by": ["Q1", "Q2", "Q3"],
            },
            {
                "path": "paper_output/data_cleaned/annual_10min.csv",
                "rows": int(len(annual_table)),
                "fields": list(annual_table.columns),
                "used_by": ["Q2", "Q3", "Q4-2", "Q4-3"],
            },
            {
                "path": "paper_output/data_cleaned/pv_forecast_long.csv",
                "rows": int(len(forecast_long)),
                "fields": list(forecast_long.columns),
                "used_by": ["Q3", "Q4-3"],
            },
            {
                "path": "paper_output/data_cleaned/pv_forecast_10min_by_origin.csv",
                "rows": int(len(forecast_by_origin)),
                "fields": list(forecast_by_origin.columns),
                "used_by": ["Q3", "Q4-3"],
            },
            {
                "path": "paper_output/data_cleaned/pv_forecast_10min_asof.csv",
                "rows": int(len(forecast_asof)),
                "fields": list(forecast_asof.columns),
                "used_by": ["Q3", "Q4-3"],
            },
        ],
        "question_links": [
            {
                "question_id": "Q1",
                "required_dataset": "attachment1_10min.csv",
                "required_fields": ["price_fixed_yuan_per_kwh", "load_kw", "pv_forecast_kw"],
            },
            {
                "question_id": "Q2",
                "required_dataset": "annual_10min.csv",
                "required_fields": ["date", "interval_id", "load_kw", "pv_actual_kw", "price_fixed_yuan_per_kwh"],
            },
            {
                "question_id": "Q3",
                "required_dataset": "annual_10min.csv + pv_forecast_10min_by_origin.csv + pv_forecast_10min_asof.csv",
                "required_fields": ["load_kw", "pv_actual_kw", "forecast_origin_hour", "forecast_horizon_hour", "pv_forecast_kw"],
            },
            {
                "question_id": "Q4",
                "required_dataset": "annual_10min.csv + pv_forecast_10min_by_origin.csv + pv_forecast_10min_asof.csv",
                "required_fields": ["price_actual_yuan_per_kwh", "load_kw", "pv_actual_kw", "pv_forecast_kw"],
            },
        ],
        "cleaning_rules": [
            "附件 3 合并单元格造成的空日期按上一行日期向下填充。",
            "附件 1、2、4 的所有功率字段乘以 1/6 转为 kWh。",
            "附件时间标签按结果模板解释为时段起点，首段 00:10-00:20，末段次日 00:00-00:10。",
            "附件 3 的整点预报按目标小时映射到 6 个 10 分钟时段。",
            "每个预报时刻保留可用和不可用状态，已经发生的时段标记 past_locked，不与未来预报混用。",
            "0:00 首小时缺少当日预报时，夜间光伏按 0 回退并明确记录状态，不生成非零伪值。",
        ],
        "template_contract": templates,
        "downstream_requirements": {
            "S4_must_read": [
                "paper_output/data_cleaned/annual_10min.csv",
                "paper_output/data_cleaned/pv_forecast_10min_by_origin.csv",
                "paper_output/data_cleaned/pv_forecast_10min_asof.csv",
            ],
            "S4_must_write": [
                "result1.xlsx",
                "result2.xlsx",
                "result3.xlsx",
                "result4-2.xlsx",
                "result4-3.xlsx",
            ],
            "S5_must_validate": [
                "功率平衡",
                "SOC 边界与跨日连续性",
                "144 时段完整性",
                "334 天输出完整性",
                "调整费用和紧急购电费用复算",
            ],
        },
    }


def build_visualization_plan(model_route, ready_figures):
    planned = build_planned_figures(model_route)
    return {
        "schema_version": "1.1",
        "stage": "S3",
        "status": "complete",
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "generated_by": "paper_output/code/data_processing/build_s3_data.py",
        "figure_count": len(planned) + len(ready_figures),
        "ready_count": len(ready_figures),
        "planned_count": len(planned),
        "figures": ready_figures + planned,
        "notes": [
            "data_cleaned 图表已经由附件原始数据生成。",
            "模型结果图在 result*.xlsx 和模型结果生成后才可由 S5 转为 ready。",
            "placeholder 或 missing 图不得进入正式论文证据。",
        ],
    }


def build_figure_index(visualization_plan):
    figures = []
    for item in visualization_plan["figures"]:
        path_text = str(item.get("output_path") or "").replace("\\", "/")
        exists = bool(path_text and (ROOT / path_text).exists())
        figures.append(
            {
                "figure_id": item.get("figure_id"),
                "question_id": item.get("question_id"),
                "title": item.get("title"),
                "path": path_text,
                "exists": exists,
                "ok": exists if item.get("status") == "ready" else False,
                "placeholder": False,
                "status": item.get("status"),
                "paper_usage": item.get("paper_usage"),
            }
        )
    return {
        "schema_version": "1.1",
        "stage": "S3",
        "status": "complete",
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "figures": figures,
        "summary": {
            "total": len(figures),
            "ready": sum(1 for item in figures if item["status"] == "ready" and item["exists"]),
            "planned": sum(1 for item in figures if item["status"] == "planned"),
        },
    }


def build_load_report(
    annual_table,
    attachment1_table,
    forecast_long,
    forecast_by_origin,
    forecast_asof,
    date_blank_count,
    monthly_stats,
    forecast_availability,
    outputs,
    templates,
):
    raw_report = load_json(RAW_LOAD_REPORT_FILE) if RAW_LOAD_REPORT_FILE.exists() else {}
    checks = [
        {
            "id": "D01",
            "name": "全年日期覆盖",
            "status": "PASS" if annual_table["date"].nunique() == 365 else "FAIL",
            "details": "{} days".format(int(annual_table["date"].nunique())),
        },
        {
            "id": "D02",
            "name": "每日 144 时段",
            "status": "PASS"
            if (annual_table.groupby("date").size() == 144).all()
            else "FAIL",
            "details": "every date must contain intervals 1..144",
        },
        {
            "id": "D03",
            "name": "输出期 334 天",
            "status": "PASS"
            if annual_table.loc[annual_table["is_official_output_period"], "date"].nunique() == 334
            else "FAIL",
            "details": "2025-02-01 to 2025-12-31",
        },
        {
            "id": "D04",
            "name": "附件 1 完整性",
            "status": "PASS" if len(attachment1_table) == 144 else "FAIL",
            "details": "{} rows".format(len(attachment1_table)),
        },
        {
            "id": "D05",
            "name": "附件 3 日期向下填充",
            "status": "PASS" if date_blank_count == 1095 else "FAIL",
            "details": "{} blank date cells filled forward".format(date_blank_count),
        },
        {
            "id": "D06",
            "name": "预报整点数据完整性",
            "status": "PASS" if len(forecast_long) == 365 * 4 * 24 else "FAIL",
            "details": "{} rows".format(len(forecast_long)),
        },
        {
            "id": "D07",
            "name": "原始数值非负",
            "status": "PASS"
            if min(
                annual_table["load_kw"].min(),
                annual_table["pv_actual_kw"].min(),
                annual_table["price_actual_yuan_per_kwh"].min(),
                attachment1_table["price_fixed_yuan_per_kwh"].min(),
                attachment1_table["pv_forecast_kw"].min(),
            )
            >= 0
            else "FAIL",
            "details": "load, PV and price must be non-negative",
        },
        {
            "id": "D08",
            "name": "按预报时刻展开",
            "status": "PASS" if len(forecast_by_origin) == 365 * 4 * 144 else "FAIL",
            "details": "{} rows".format(len(forecast_by_origin)),
        },
        {
            "id": "D09",
            "name": "最新可用预报映射",
            "status": "PASS" if len(forecast_asof) == 365 * 144 else "FAIL",
            "details": "{} rows".format(len(forecast_asof)),
        },
        {
            "id": "D10",
            "name": "结果模板可读取",
            "status": "PASS" if len(templates) == 5 else "FAIL",
            "details": "{} templates".format(len(templates)),
        },
    ]
    unexpected_missing = {}
    asof_allowed_missing = {
        "forecast_source_origin_hour",
        "forecast_horizon_hour",
        "source_column",
    }
    for frame, name in [
        (annual_table, "annual_10min"),
        (attachment1_table, "attachment1_10min"),
        (forecast_long, "pv_forecast_long"),
        (forecast_asof, "pv_forecast_10min_asof"),
    ]:
        missing_columns = {
            column: int(count)
            for column, count in frame.isna().sum().items()
            if int(count) > 0
        }
        if name == "pv_forecast_10min_asof":
            missing_columns = {
                column: count
                for column, count in missing_columns.items()
                if column not in asof_allowed_missing
            }
        if missing_columns:
            unexpected_missing[name] = missing_columns

    expected_missing = {
        "pv_forecast_10min_by_origin": {
            "pv_forecast_kw": int(forecast_by_origin["pv_forecast_kw"].isna().sum()),
            "pv_forecast_kwh": int(forecast_by_origin["pv_forecast_kwh"].isna().sum()),
            "reason": "past_locked 或 current_hour_hold 状态不填入可用预报值。",
        },
        "pv_forecast_10min_asof": {
            "forecast_source_origin_hour": int(
                forecast_asof["forecast_source_origin_hour"].isna().sum()
            ),
            "forecast_horizon_hour": int(
                forecast_asof["forecast_horizon_hour"].isna().sum()
            ),
            "source_column": int(forecast_asof["source_column"].eq("").sum()),
            "reason": "数据集首日 00:10 前缺少前一日 18:00 预报，按夜间零光伏回退。",
        },
    }

    return {
        "schema_version": "1.1",
        "stage": "S3",
        "status": "PASS" if all(item["status"] == "PASS" for item in checks) else "FAIL",
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "generated_by": "paper_output/code/data_processing/build_s3_data.py",
        "input_manifest": "paper_output/input_manifest.json",
        "input_manifest_used": True,
        "input_manifest_sha256": sha256_file(MANIFEST_FILE),
        "raw_loader_report": {
            "status": raw_report.get("status"),
            "generated_at": raw_report.get("generated_at"),
            "data_file_count": raw_report.get("summary", {}).get("data_file_count"),
        },
        "inputs": {
            "附件1": {
                "path": "problem_files/附件1.xlsx",
                "sha256": sha256_file(ATTACHMENT1),
                "rows": len(attachment1_table),
                "columns": list(attachment1_table.columns),
            },
            "附件2": {
                "path": "problem_files/附件2.xlsx",
                "sha256": sha256_file(ATTACHMENT2),
                "sheets": ["小区负载", "光伏发电实际功率"],
                "annual_rows": int(len(annual_table)),
            },
            "附件3": {
                "path": "problem_files/附件3.xlsx",
                "sha256": sha256_file(ATTACHMENT3),
                "raw_rows": 1460,
                "blank_date_cells_forward_filled": date_blank_count,
                "forecast_long_rows": int(len(forecast_long)),
            },
            "附件4": {
                "path": "problem_files/附件4.xlsx",
                "sha256": sha256_file(ATTACHMENT4),
                "annual_rows": 365,
            },
        },
        "cleaned_outputs": outputs,
        "profiles": {
            "annual_10min": {
                "load_kw": numeric_summary(annual_table["load_kw"]),
                "pv_actual_kw": numeric_summary(annual_table["pv_actual_kw"]),
                "price_actual_yuan_per_kwh": numeric_summary(
                    annual_table["price_actual_yuan_per_kwh"]
                ),
            },
            "attachment1_10min": {
                "price_fixed_yuan_per_kwh": numeric_summary(
                    attachment1_table["price_fixed_yuan_per_kwh"]
                ),
                "load_kw": numeric_summary(attachment1_table["load_kw"]),
                "pv_forecast_kw": numeric_summary(attachment1_table["pv_forecast_kw"]),
            },
        },
        "mapping_contract": {
            "interval_semantics": "attachment time label is interval start",
            "first_slot": "00:10-00:20",
            "last_slot": "next day 00:00-00:10",
            "forecast_expansion": "integer-hour forecast repeated across six 10-minute slots",
            "asof_origin_rule": "use the latest available origin strictly before the slot start with a positive forecast horizon",
            "past_data_protection": "past_locked rows are excluded from usable forecast values",
        },
        "checks": checks,
        "expected_missing": expected_missing,
        "unexpected_missing": unexpected_missing,
        "monthly_stats": monthly_stats.to_dict(orient="records"),
        "forecast_availability": forecast_availability.to_dict(orient="records"),
        "templates": templates,
        "warnings": [
            "result2/result3 的充放电量和紧急购电工作表是模板样式，正式输出时需要扩展至 334 天。",
            "附件 3 的首日 00:00 前没有前一天 18:00 预报，相关 10 分钟槽使用夜间零光伏回退。",
        ],
    }


def main():
    configure_plot_style()
    CLEAN_DIR.mkdir(parents=True, exist_ok=True)
    PLAN_DIR.mkdir(parents=True, exist_ok=True)
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    EDA_FIGURE_DIR.mkdir(parents=True, exist_ok=True)

    attachment1_header, attachment1_rows = load_attachment1()
    load_rows, pv_rows = load_attachment2()
    forecast_rows = load_attachment3()
    price_rows = load_attachment4()
    templates = load_result_templates()

    result1_labels = result1_interval_labels()
    result2_labels = result2_interval_labels()
    interval_metadata = build_interval_metadata(attachment1_rows, result1_labels)

    attachment1_table = build_attachment1_table(
        attachment1_rows,
        interval_metadata,
        result2_labels,
    )
    annual_table = build_annual_table(
        load_rows,
        pv_rows,
        price_rows,
        attachment1_table,
        interval_metadata,
        result2_labels,
    )
    forecast_long, date_blank_count = build_forecast_long(forecast_rows)
    forecast_by_origin = build_forecast_by_origin(
        annual_table,
        forecast_long,
    )
    forecast_asof = build_forecast_asof(
        annual_table,
        forecast_long,
    )

    monthly_stats = build_monthly_stats(annual_table)
    forecast_availability = build_forecast_availability(
        forecast_by_origin,
        forecast_asof,
    )

    output_paths = {
        "attachment1_10min": CLEAN_DIR / "attachment1_10min.csv",
        "annual_10min": CLEAN_DIR / "annual_10min.csv",
        "pv_forecast_long": CLEAN_DIR / "pv_forecast_long.csv",
        "pv_forecast_10min_by_origin": CLEAN_DIR / "pv_forecast_10min_by_origin.csv",
        "pv_forecast_10min_asof": CLEAN_DIR / "pv_forecast_10min_asof.csv",
        "monthly_stats": CLEAN_DIR / "monthly_stats.csv",
        "forecast_availability": CLEAN_DIR / "forecast_availability.csv",
    }
    attachment1_table.to_csv(output_paths["attachment1_10min"], index=False, encoding="utf-8-sig")
    annual_table.to_csv(output_paths["annual_10min"], index=False, encoding="utf-8-sig")
    forecast_long.to_csv(output_paths["pv_forecast_long"], index=False, encoding="utf-8-sig")
    forecast_by_origin.to_csv(
        output_paths["pv_forecast_10min_by_origin"],
        index=False,
        encoding="utf-8-sig",
    )
    forecast_asof.to_csv(
        output_paths["pv_forecast_10min_asof"],
        index=False,
        encoding="utf-8-sig",
    )
    monthly_stats.to_csv(output_paths["monthly_stats"], index=False, encoding="utf-8-sig")
    forecast_availability.to_csv(
        output_paths["forecast_availability"],
        index=False,
        encoding="utf-8-sig",
    )

    ready_figures = write_eda_figures(
        annual_table,
        forecast_asof,
        monthly_stats,
    )
    model_route = load_json(MODEL_ROUTE_FILE)
    data_plan = build_data_plan(
        annual_table,
        attachment1_table,
        forecast_long,
        forecast_by_origin,
        forecast_asof,
        templates,
    )
    visualization_plan = build_visualization_plan(model_route, ready_figures)
    figure_index = build_figure_index(visualization_plan)

    output_profiles = []
    for path in output_paths.values():
        frame = pd.read_csv(path, encoding="utf-8-sig")
        output_profiles.append(dataframe_profile(frame, path))

    load_report = build_load_report(
        annual_table,
        attachment1_table,
        forecast_long,
        forecast_by_origin,
        forecast_asof,
        date_blank_count,
        monthly_stats,
        forecast_availability,
        output_profiles,
        templates,
    )

    write_json(PLAN_DIR / "data_plan.json", data_plan)
    write_json(PLAN_DIR / "visualization_plan.json", visualization_plan)
    write_json(OUTPUT_DIR / "figure_index.json", figure_index)
    write_json(CLEAN_DIR / "load_report.json", load_report)

    print("S3 data cleaning complete")
    print("load_report: {}".format((CLEAN_DIR / "load_report.json").relative_to(ROOT)))
    print("data_plan: {}".format((PLAN_DIR / "data_plan.json").relative_to(ROOT)))
    print("visualization_plan: {}".format((PLAN_DIR / "visualization_plan.json").relative_to(ROOT)))
    print("figure_index: {}".format((OUTPUT_DIR / "figure_index.json").relative_to(ROOT)))


if __name__ == "__main__":
    main()
